"""生成Excel、Word和Markdown转写成品，并回读核验Office文件。"""

import html
import math
import os
import re
import string
import tempfile
import unicodedata
from datetime import timedelta
from pathlib import Path
from typing import Protocol, cast
from zipfile import BadZipFile

from docx import Document
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from ..models import Sentence, Transcript

FONT_NAME = "等线"
SHEET = "转写明细"
HEADERS = ("序号", "音轨", "开始时间", "结束时间", "说话人", "转写内容")
FIRST_ROW = 4
TIME_FORMAT = "[hh]:mm:ss.000"
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


class DocumentError(ValueError):
    """表示可向用户展示的文档格式或写入错误。"""


class DocumentWriter(Protocol):
    """约定三种文档生成函数共用的输入参数。"""

    def __call__(self, transcript: Transcript, path: Path, *, source_name: str, job_id: str, model: str) -> None:
        """将转写内容与任务标签写入指定格式文件。"""
        ...


def timestamp(milliseconds: int) -> str:
    """将毫秒转换为累计小时的时分秒文本。"""
    seconds, millis = divmod(milliseconds, 1000)
    minutes, seconds = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}.{millis:03d}"


def publish_document(writer: DocumentWriter, transcript: Transcript, final_path: Path, *,
                     source_name: str, job_id: str, model: str) -> int:
    """先生成临时文件，再替换同名成品并返回文件大小。"""
    with tempfile.NamedTemporaryFile(dir=final_path.parent,
                                      prefix=f"{final_path.stem}.partial-pid-{os.getpid()}-",
                                      suffix=final_path.suffix, delete=False) as stream:
        temporary = Path(stream.name)
    # Windows需先关闭临时句柄；每次发布只写入和清理自己创建的文件。
    try:
        writer(transcript, temporary, source_name=source_name, job_id=job_id, model=model)
        temporary.replace(final_path)
        return final_path.stat().st_size
    finally:
        temporary.unlink(missing_ok=True)


def _title(source_name: str) -> str:
    """生成三种成品共用的“源文件名 录音转写”标题。"""
    return f"{Path(source_name).stem} 录音转写"


def _speaker(sentence: Sentence) -> str:
    """保留原始说话人编号，缺失时明确标为未标注。"""
    return "未标注" if sentence.speaker_id is None else f"说话人{sentence.speaker_id}"


def _track(sentence: Sentence) -> str:
    """生成音轨标签；缺少音轨号时保留结果组序号。"""
    if sentence.channel_id is None:
        return f"未标注（第{sentence.track_index}组）"
    return f"音轨{sentence.channel_id}"


def _label(sentence: Sentence, *, separator: str) -> str:
    """组合段落序号、时间范围、说话人与音轨供文本成品使用。"""
    return separator.join((str(sentence.index),
                           f"{timestamp(sentence.begin_ms)} — {timestamp(sentence.end_ms)}",
                           _speaker(sentence), _track(sentence)))


def _xml_text(text: str, format_name: str) -> None:
    """检查Office正文字符是否符合XML存储要求。"""
    if INVALID_XML.search(text):
        raise DocumentError(f"{format_name}无法保存结果中的控制字符；原始JSON及其他格式不受影响。")


def _cell_value(cell: Cell, value: str | int | timedelta) -> None:
    """写入Excel单元格，检查文本存储上限并固定字符串类型。"""
    if isinstance(value, str):
        _xml_text(value, "Excel")
        if len(value) > 32767:
            raise DocumentError("Excel单元格最多保存32767个字符，本次未截断内容。请查看JSON或其他成品。")
        cell.value = value
        # 转写和文件名都是不可信文本；即使以=开头也不能成为Excel公式。
        cell.data_type = "s"
    else:
        cell.value = value


def _xlsx_values(sentence: Sentence) -> tuple[int, str, timedelta, timedelta, str, str]:
    """生成与Excel表头顺序一致的段落字段。"""
    return (sentence.index, _track(sentence), timedelta(milliseconds=sentence.begin_ms),
            timedelta(milliseconds=sentence.end_ms), _speaker(sentence), sentence.text)


def _row_height(text: str, width: float) -> float:
    """估算正文展示高度，并限制在Excel行高上限内。"""
    lines = 0
    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        units = sum(4 if char == "\t" else 2.2 if unicodedata.east_asian_width(char) in "WF"
                    else 1.1 for char in line)
        lines += max(1, math.ceil(units / (width - 4)))
    return min(409, max(30, 20 * lines + 10))


def write_xlsx(transcript: Transcript, path: Path, *, source_name: str, job_id: str, model: str) -> None:
    """写入带时间、标签和正文的Excel，并回读核验内容。"""
    if len(transcript.sentences) + FIRST_ROW - 1 > 1_048_576:
        raise DocumentError("转写段落超过Excel工作表行数上限，本次未截断内容。")
    workbook = Workbook()
    try:
        # 新建工作簿包含一个普通工作表。
        sheet = cast(Worksheet, workbook.active)
        sheet.title = SHEET
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 85
        widths = (8, 22, 18, 18, 18, 88)
        for column_letter, width in zip("ABCDEF", widths):
            sheet.column_dimensions[column_letter].width = width
        metadata = (_title(source_name), f"模型：{model}    任务：{job_id}")
        workbook.properties.title = metadata[0]
        for row, value in enumerate(metadata, 1):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            cell = cast(Cell, sheet.cell(row, 1))
            _cell_value(cell, value)
            cell.font = Font(name=FONT_NAME, size=16 if row == 1 else 10,
                             bold=row == 1, color="000000")
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
            sheet.row_dimensions[row].height = 38 if row == 1 else 30
        for column, value in enumerate(HEADERS, 1):
            cell = cast(Cell, sheet.cell(FIRST_ROW - 1, column))
            _cell_value(cell, value)
            cell.font = Font(name=FONT_NAME, size=11, bold=True, color="000000")
            cell.border = Border(bottom=Side(style="thin", color="000000"))
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        sheet.row_dimensions[FIRST_ROW - 1].height = 30
        body_font = Font(name=FONT_NAME, size=11, color="000000")
        bottom = Border(bottom=Side(style="hair", color="000000"))
        for row, sentence in enumerate(transcript.sentences, FIRST_ROW):
            values: tuple[str | int | timedelta, ...] = _xlsx_values(sentence)
            for column, body_value in enumerate(values, 1):
                cell = cast(Cell, sheet.cell(row, column))
                _cell_value(cell, body_value)
                cell.font = body_font
                cell.border = bottom
                cell.alignment = Alignment(horizontal="left" if column == 6 else "center",
                                           vertical="center",
                                           wrap_text=column == 6)
                if column in (3, 4):
                    cell.number_format = TIME_FORMAT
            sheet.row_dimensions[row].height = _row_height(sentence.text, widths[-1])
        sheet.freeze_panes = f"B{FIRST_ROW}"
        sheet.auto_filter.ref = f"A{FIRST_ROW - 1}:F{sheet.max_row}"
        sheet.print_title_rows = f"1:{FIRST_ROW - 1}"
        sheet.print_area = f"A1:F{sheet.max_row}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        cast(PageSetupProperties, sheet.sheet_properties.pageSetUpPr).fitToPage = True
        workbook.save(path)
    except OSError as exc:
        raise DocumentError("Excel写入失败，请检查保存位置和文件占用；未自动重试。") from exc
    finally:
        workbook.close()
    try:
        # 使用上下文管理器，在回读结束后关闭文件句柄。
        with path.open("rb") as stream:
            verified = load_workbook(stream, read_only=True, data_only=False)
            try:
                restored_sheet = verified[SHEET]
                if (restored_sheet.max_row != len(transcript.sentences) + FIRST_ROW - 1
                        or restored_sheet.max_column != 6 or restored_sheet["A1"].value != metadata[0]
                        or restored_sheet["A2"].value != metadata[1]):
                    raise DocumentError("Excel回读的段落数或元信息不一致，未报告完成。")
                rows = restored_sheet.iter_rows(min_row=FIRST_ROW, max_col=6)
                for cells, sentence in zip(rows, transcript.sentences, strict=True):
                    restored_values = tuple(cell.value if cell.value is not None else "" for cell in cells)
                    if (restored_values != _xlsx_values(sentence)
                            or any(cell.data_type == "f" for cell in cells)):
                        raise DocumentError("Excel回读的正文、时间或标签不一致，未报告完成。")
            finally:
                verified.close()
    except (OSError, ValueError, KeyError, BadZipFile) as exc:
        if isinstance(exc, DocumentError):
            raise
        raise DocumentError("Excel无法重新打开核验，未报告完成。") from exc


def _word_text(paragraph: Paragraph, text: str) -> None:
    """写入Word正文，并保留python-docx默认会转换的回车字符。"""
    _xml_text(text, "Word")
    # python-docx把CR转换成LF。独立写入CR文本节点，以便回读时仍保留原始字符。
    for index, part in enumerate(text.split("\r")):
        if index:
            element = OxmlElement("w:t")
            element.text = "\r"
            paragraph.add_run()._r.append(element)
        paragraph.add_run(part)


def _word_style(style: ParagraphStyle, *, size: float, bold: bool = False) -> None:
    """统一Word字体与字号，清除模板主题对字体选择的覆盖。"""
    style.font.name = FONT_NAME
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.cs_bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    properties = style.element.get_or_add_rPr()
    # Word对阿拉伯语等复杂文字使用独立的半磅字号，覆盖模板中的默认值。
    complex_size = properties.find(qn("w:szCs"))
    if complex_size is None:
        complex_size = OxmlElement("w:szCs")
        properties.get_or_add_sz().addnext(complex_size)
    complex_size.set(qn("w:val"), str(int(size * 2)))
    fonts = properties.rFonts
    # 同时覆盖中西文与复杂文字字体，移除模板主题字体对当前样式的影响。
    for script in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{script}"), FONT_NAME)
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attribute}"), None)


def write_docx(transcript: Transcript, path: Path, *, source_name: str, job_id: str, model: str) -> None:
    """写入按段落排版的Word，并回读核验内容。"""
    expected = [_title(source_name), f"模型：{model}", f"任务：{job_id}"]
    try:
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.85)
        normal_style = cast(ParagraphStyle, document.styles["Normal"])
        title_style = cast(ParagraphStyle, document.styles["Title"])
        footer_style = cast(ParagraphStyle, document.styles["Footer"])
        _word_style(normal_style, size=10)
        body_format = normal_style.paragraph_format
        body_format.line_spacing = 1.25
        body_format.space_after = Pt(6)
        body_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        body_format.keep_together = False
        body_format.widow_control = True
        _word_style(title_style, size=20, bold=True)
        title_style.paragraph_format.space_after = Pt(12)
        title_style.paragraph_format.keep_with_next = True
        metadata_style = cast(ParagraphStyle, document.styles.add_style("Transcript Metadata", WD_STYLE_TYPE.PARAGRAPH))
        metadata_style.base_style = normal_style
        _word_style(metadata_style, size=10)
        metadata_style.paragraph_format.space_after = Pt(4)
        metadata_style.paragraph_format.keep_with_next = True
        marker_style = cast(ParagraphStyle, document.styles.add_style("Transcript Marker", WD_STYLE_TYPE.PARAGRAPH))
        marker_style.base_style = normal_style
        _word_style(marker_style, size=10, bold=True)
        marker_style.paragraph_format.keep_with_next = True
        marker_style.paragraph_format.space_before = Pt(6)
        marker_style.paragraph_format.space_after = Pt(3)
        for index, text in enumerate(expected):
            paragraph = document.add_paragraph(style="Title" if index == 0 else metadata_style)
            _word_text(paragraph, text)
        for sentence in transcript.sentences:
            label = _label(sentence, separator="  ")
            _word_text(document.add_paragraph(style=marker_style), label)
            _word_text(document.add_paragraph(), sentence.text)
            expected.extend((label, sentence.text))
        footer = section.footer.paragraphs[0]
        _word_style(footer_style, size=10)
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("第 ")
        page = OxmlElement("w:fldSimple")
        page.set(qn("w:instr"), "PAGE")
        footer._p.append(page)
        footer.add_run(" 页")
        document.core_properties.author = "asr-transcription"
        document.save(str(path))
        with path.open("rb") as stream:
            restored = Document(stream)
            if [paragraph.text for paragraph in restored.paragraphs] != expected:
                raise DocumentError("Word回读的正文、时间或标签不一致，未报告完成。")
    except OSError as exc:
        raise DocumentError("Word写入或回读失败，请检查保存位置和文件占用；未自动重试。") from exc


def _markdown_text(text: str) -> str:
    """转义正文中的Markdown和HTML语法，保留文字及换行。"""
    # CommonMark反斜杠转义阻止链接/图像等语法；HTML字符另行编码。
    # 只编码行首空白来避免代码块，英文正文保留普通空格，便于直接阅读源文件。
    result = []
    line_start = True
    for char in text:
        if char == "\n":
            result.append(char)
            line_start = True
            continue
        if (line_start and char in " \t") or char == "\r":
            result.append(f"&#{ord(char)};")
            continue
        line_start = False
        if char in "<>&":
            result.append(html.escape(char))
        elif char in string.punctuation:
            result.append("\\" + char)
        else:
            result.append(char)
    return "".join(result)


def write_markdown(transcript: Transcript, path: Path, *, source_name: str, job_id: str, model: str) -> None:
    """写入按段落组织的UTF-8 Markdown，保留时间、标签与原文字符。"""
    parts = [f"# {_markdown_text(_title(source_name))}", f"模型：{_markdown_text(model)}",
             f"任务：{_markdown_text(job_id)}"]
    for sentence in transcript.sentences:
        parts.extend((f"### {_label(sentence, separator=' · ')}", _markdown_text(sentence.text)))
    content = "\n\n".join(parts) + "\n"
    try:
        with path.open("w", encoding="utf-8", newline="") as output:
            output.write(content)
    except OSError as exc:
        raise DocumentError("Markdown写入失败，请检查保存位置和文件占用；未自动重试。") from exc
