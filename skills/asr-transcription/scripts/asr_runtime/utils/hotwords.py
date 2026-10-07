"""读取热词Excel原始行并生成词表模板。"""

import io
import math
import zipfile
from typing import cast
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from ..models import MAX_HOTWORD_ROWS, HotwordField, HotwordRow, HotwordValue, LocalizedText
from .i18n import localize, translate


# 本机工作簿解析资源上限，不是阿里云接口限制。
MAX_XLSX_BYTES = 5_000_000
MAX_XLSX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_XLSX_ENTRIES = 200


class HotwordFileError(ValueError):
    """表示热词Excel文件结构错误。"""

    def __init__(self, message: str) -> None:
        """保留文件错误原文，并提供当前语言的异常说明。"""
        super().__init__(translate(message))
        self.template = message


def _check_xlsx_archive(stream: io.BytesIO) -> None:
    """检查内存中Excel压缩包的展开规模和加密标记。"""
    with zipfile.ZipFile(stream) as archive:
        entries = archive.infolist()
        if (len(entries) > MAX_XLSX_ENTRIES
                or sum(item.file_size for item in entries) > MAX_XLSX_UNCOMPRESSED_BYTES):
            raise HotwordFileError("热词Excel结构过大：解压内容最多20 MiB、200个内部文件。")
        if any(item.flag_bits & 1 for item in entries):
            raise HotwordFileError("不支持加密的热词Excel，请保存为普通.xlsx文件。")


def read_hotwords(content: bytes) -> tuple[list[HotwordRow], list[LocalizedText]]:
    """从Excel字节读取固定两列原始值及工作簿提示。"""
    workbook = None
    try:
        if len(content) > MAX_XLSX_BYTES:
            raise HotwordFileError("热词文件超过 5 MB 文件上限。")
        # 有界的小型工作簿直接读取，不依赖可伪造的dimension标签。
        with io.BytesIO(content) as stream:
            _check_xlsx_archive(stream)
            stream.seek(0)
            workbook = load_workbook(stream, read_only=False, data_only=False, keep_links=False)
        sheet: object
        if "热词" in workbook.sheetnames:
            sheet = workbook["热词"]
        elif len(workbook.sheetnames) == 1:
            sheet = workbook.active
        else:
            raise HotwordFileError("多个工作表时请将待使用的工作表命名为“热词”。")
        if not isinstance(sheet, Worksheet):
            raise HotwordFileError("请使用普通工作表填写热词，不支持图表工作表。")
        if sheet.max_row > MAX_HOTWORD_ROWS + 1 or sheet.max_column > 2:
            raise HotwordFileError("热词工作表仅支持两列、最多10001行（含表头和空行）。")
        header = [sheet.cell(1, number).value for number in (1, 2)]
        if header[0] not in ("text", "热词") or header[1] not in ("weight", "权重"):
            raise HotwordFileError("当前 Excel 未按模板导入，请下载模板，按模板填写后重新导入。")
        warnings = []
        if len(workbook.sheetnames) > 1:
            warnings.append(localize("仅读取名为“热词”的工作表，其他工作表不参与此次转写。"))
        rows: list[HotwordRow] = []
        for cells in sheet.iter_rows(min_row=2, max_col=2):
            row: HotwordRow = {"text": None, "weight": None}
            invalid_fields: list[HotwordField] = []
            for field, cell in zip(("text", "weight"), cells):
                name = cast(HotwordField, field)
                value = cell.value
                # 公式、日期和Excel错误值保留可见内容，用户改写对应单元格后移除此类型标记。
                if (cell.data_type in ("e", "f") or value is not None and not isinstance(value, (str, int, float, bool))
                        or isinstance(value, float) and not math.isfinite(value)):
                    row[name] = str(value)
                    invalid_fields.append(name)
                else:
                    row[name] = cast(HotwordValue, value)
            if invalid_fields:
                row["invalid_fields"] = invalid_fields
            if row["text"] in (None, "") and row["weight"] in (None, "") and not invalid_fields:
                continue
            rows.append(row)
        return rows, warnings
    except (OSError, zipfile.BadZipFile, InvalidFileException, ParseError,
            DefusedXmlException, KeyError, ValueError) as exc:
        if isinstance(exc, HotwordFileError):
            raise
        raise HotwordFileError("无法读取热词Excel，请检查文件是否损坏或仍在写入。") from exc
    finally:
        if workbook is not None:
            workbook.close()


def hotwords_template() -> bytes:
    """生成带示例词条的两列Excel模板，供网页直接下载。"""
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "热词"
    sheet.append(["text", "weight"])
    sheet.append([translate("示例术语"), 4])
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 12
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
