"""验证可编辑热词表格和上下文错误的定位与修正。"""

import json
from datetime import datetime
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from asr_runtime.application.inputs import import_hotwords
from asr_runtime.application.rules import ValidationError, validate_context, validate_hotword_rows
from asr_runtime.models import MAX_HOTWORD_ROWS
from asr_runtime.utils.i18n import language_scope
from tests.support import RuntimeTestCase


class HotwordEditorTests(RuntimeTestCase):
    def test_duplicate_words_mark_every_row_even_when_weights_are_invalid(self):
        """验证整组重复词都标错，非法权重同时保留独立提示。"""
        rows = [{"text": "IPO", "weight": weight} for weight in (4, -1000, 3, 4)]
        with self.assertRaises(ValidationError) as caught:
            validate_hotword_rows(rows)
        details = caught.exception.details
        self.assertEqual([item["row"] for item in details if item["field"] == "text"], [1, 2, 3, 4])
        self.assertEqual([item["row"] for item in details if item["field"] == "weight"], [2])
        self.assertTrue(all("重复" in item["message"] for item in details if item["field"] == "text"))
        self.assertEqual(validate_hotword_rows(rows[:1])["vocabulary"], {"IPO": 4})

    def workbook(self, rows):
        """保存带固定表头的测试工作簿。"""
        path = self.runtime.path("editor.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "热词"
        sheet.append(["text", "weight"])
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def test_import_preserves_invalid_rows_and_edits_replace_file_input(self):
        """验证导入保留原值，词条问题在明确校验时报告。"""
        path = self.workbook([("术语", 4), ("另一个词", "错误"), ("术语", 5)])
        original = path.read_bytes()
        with patch("asr_runtime.application.rules.build_vocabulary", side_effect=AssertionError("导入只读取原始行")):
            report = import_hotwords(path.read_bytes())
        self.assertNotIn("issues", report)
        with self.assertRaises(ValidationError) as caught:
            validate_hotword_rows(report["rows"])
        self.assertEqual([(issue["row"], issue["field"]) for issue in caught.exception.details],
                         [(1, "text"), (2, "weight"), (3, "text")])
        self.assertEqual(report["rows"][1]["weight"], "错误")
        self.assertEqual(path.read_bytes(), original)

        path.unlink()
        report["rows"][1]["weight"] = "4"
        report["rows"].pop(2)
        result = validate_hotword_rows(report["rows"])
        self.assertEqual(result["vocabulary"], {"术语": 4, "另一个词": 4})
        self.assertEqual(result["warnings"], [])

    def test_formula_errors_belong_to_cells_and_disappear_after_edit(self):
        """验证公式原文及错误定位可见，改写后使用当前单元格值。"""
        report = import_hotwords(self.workbook([("=1+1", 4), ("另一个词", "=2+2")]).read_bytes())
        self.assertEqual(report["rows"], [
            {"text": "=1+1", "weight": 4, "invalid_fields": ["text"]},
            {"text": "另一个词", "weight": "=2+2", "invalid_fields": ["weight"]},
        ])
        self.assertNotIn("issues", report)
        with self.assertRaises(ValidationError) as caught:
            validate_hotword_rows(report["rows"])
        self.assertEqual([(issue["row"], issue["field"]) for issue in caught.exception.details],
                         [(1, "text"), (2, "weight")])
        self.assertTrue(all("不接受公式" in issue["message"] for issue in caught.exception.details))
        report["rows"][0].update(text="术语", invalid_fields=[])
        report["rows"][1].update(weight="4", invalid_fields=[])
        self.assertEqual(validate_hotword_rows(report["rows"])["count"], 2)

    def test_literal_equals_prefix_is_distinct_from_excel_formula(self):
        """验证相同的等号前缀文本仅在Excel确认为公式时需要改写。"""
        path = self.workbook([("=1+1", 4), ("=1+1", 4)])
        workbook = load_workbook(path)
        workbook.active["A3"].data_type = "s"
        workbook.save(path)
        workbook.close()
        report = import_hotwords(path.read_bytes())
        self.assertNotIn("issues", report)
        with self.assertRaises(ValidationError) as caught:
            validate_hotword_rows(report["rows"])
        self.assertEqual([issue["row"] for issue in caught.exception.details if "公式" in issue["message"]], [1])
        self.assertEqual([issue["row"] for issue in caught.exception.details if "重复" in issue["message"]], [1, 2])
        self.assertNotIn("invalid_fields", report["rows"][1])
        self.assertEqual(validate_hotword_rows([report["rows"][1]])["vocabulary"], {"=1+1": 4})
        # 网页输入保存为固定文本，编辑动作会移除原Excel的类型标记。
        report["rows"][0]["invalid_fields"] = []
        self.assertEqual(validate_hotword_rows(report["rows"][:1])["vocabulary"], {"=1+1": 4})

    def test_limits_report_all_excess_rows_at_once(self):
        """验证总词数和超级词超限时一次返回全部超限行。"""
        for count, weight, expected in ((53, 50, [(51, "weight"), (52, "weight"), (53, "weight")]),
                                         (2002, 4, [(2001, "text"), (2002, "text")])):
            rows = [{"text": f"term{index}", "weight": weight} for index in range(count)]
            with self.subTest(count=count, weight=weight), self.assertRaises(ValidationError) as caught:
                validate_hotword_rows(rows)
            self.assertEqual([(issue["row"], issue["field"]) for issue in caught.exception.details], expected)

    def test_dates_and_excel_errors_are_visible_serializable_and_require_edit(self):
        """验证日期和Excel错误保留可见值及格级标记，编辑后可重新校验。"""
        report = import_hotwords(self.workbook([(datetime(2026, 1, 2), 4), ("另一个词", "#VALUE!")]).read_bytes())
        json.dumps(report, allow_nan=False)
        self.assertIn("2026-01-02", report["rows"][0]["text"])
        self.assertEqual(report["rows"][0]["invalid_fields"], ["text"])
        self.assertEqual(report["rows"][1]["weight"], "#VALUE!")
        self.assertEqual(report["rows"][1]["invalid_fields"], ["weight"])
        with self.assertRaises(ValidationError):
            validate_hotword_rows(report["rows"])
        report["rows"][0].update(text="术语", invalid_fields=[])
        report["rows"][1].update(weight="4", invalid_fields=[])
        self.assertEqual(validate_hotword_rows(report["rows"])["count"], 2)

    def test_text_weights_use_explicit_integer_spelling(self):
        """验证直接填写的整数权重可用且错误原值保持可修改。"""
        for weight in (1, 4.0, "1", "4", "50"):
            with self.subTest(weight=weight):
                self.assertEqual(validate_hotword_rows([{"text": "术语", "weight": weight}])["count"], 1)
        for weight in (True, False, "04", "4.0", " 4", "4 ", "5e1", "４", "", None, 6, 2.5):
            row = {"text": "术语", "weight": weight}
            with self.subTest(weight=weight), self.assertRaises(ValidationError) as caught:
                validate_hotword_rows([row])
            self.assertEqual(row["weight"], weight)
            self.assertEqual(caught.exception.details[0]["row"], 1)
            self.assertEqual(caught.exception.details[0]["field"], "weight")

    def test_empty_table_reports_a_cell_and_accepts_new_manual_row(self):
        """验证空模板保留添加词条的操作路径。"""
        report = import_hotwords(self.workbook([]).read_bytes())
        self.assertEqual(report, {"rows": [], "warnings": []})
        for rows in ([], [{"text": "", "weight": ""}]):
            with self.subTest(rows=rows), self.assertRaises(ValidationError) as caught:
                validate_hotword_rows(rows)
            self.assertEqual(caught.exception.details[0]["field"], "text")
        result = validate_hotword_rows([{"text": "Kubernetes", "weight": "4"}])
        self.assertEqual(result["vocabulary"], {"Kubernetes": 4})

    def test_row_protocol_uses_array_position_and_rejects_extra_fields(self):
        """验证当前数组就是行号来源，额外标识和非单元格内容被拒绝。"""
        valid = {"text": "术语", "weight": 4}
        for rows in (None, {}, [{**valid, "row": 1}], [{**valid, "key": "client-id"}],
                     [{**valid, "text": ["术语"]}], [{**valid, "invalid_fields": ["other"]}],
                     [valid] * (MAX_HOTWORD_ROWS + 1)):
            with self.subTest(rows_type=type(rows).__name__), self.assertRaises(ValidationError):
                validate_hotword_rows(rows)

    def test_removing_rows_repositions_errors_using_current_array(self):
        """验证删除前序词条后，剩余错误使用连续的新序号定位。"""
        rows = [{"text": f"term{index}", "weight": 4} for index in range(51)]
        rows[-1]["weight"] = "incorrect"
        with self.assertRaises(ValidationError) as before:
            validate_hotword_rows(rows)
        self.assertEqual(before.exception.details[0]["row"], 51)
        del rows[:2]
        with self.assertRaises(ValidationError) as after:
            validate_hotword_rows(rows)
        self.assertEqual(after.exception.details[0]["row"], 49)

    def test_context_messages_name_empty_length_and_character_problems(self):
        """验证参考文本错误说明具体问题并支持中英文。"""
        for language, expected in (("zh-CN", ("为空", "超出 3", "第 2", "U+0000")),
                                   ("en", ("empty", "by 3", "Character 2", "U+0000"))):
            with language_scope(language):
                for value, fragment in ((" \n", expected[0]), ("文" * 403, expected[1]),
                                        ("a\x00b", expected[2]), ("a\x00b", expected[3])):
                    with self.subTest(language=language, fragment=fragment), self.assertRaises(ValidationError) as caught:
                        validate_context(value)
                    self.assertIn(fragment, str(caught.exception))
                    self.assertEqual(caught.exception.field, "context")
        # 官方允许词表、自然语言和混合内容，本机仅检查可传输的文本规格。
        text = "Kubernetes、Bulge Bracket\n补充参考文字🙂"
        self.assertEqual(validate_context(text), text)
