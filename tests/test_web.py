import http.client
import io
import json
import threading
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import quote
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from asr_runtime.models import AudioInfo
from asr_runtime.application.transcription import job_status
from asr_runtime.web import create_server
from tests.test_session import WebFixture


class WebServerTests(WebFixture):
    def setUp(self):
        """启动本机HTTP测试服务并准备会话令牌。"""
        super().setUp()
        self.server = create_server(self.runtime)
        self.payload = self.payload_for(self.server.session)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.origin = f"http://127.0.0.1:{self.server.server_port}"

    def tearDown(self):
        """关闭本机测试服务并清理临时项目。"""
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=5)
        super().tearDown()

    def request(self, method, path, payload=None, *, token=True, headers=None, body=None):
        """向本机服务发送指定请求并解析响应。"""
        connection = http.client.HTTPConnection("127.0.0.1", self.server.server_port, timeout=5)
        supplied = {"Origin": self.origin}
        if token:
            supplied["X-ASR-Token"] = self.server.session.token
        if headers:
            supplied.update(headers)
        if payload is not None:
            body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            supplied["Content-Type"] = "application/json"
        connection.request(method, path, body=body, headers=supplied)
        response = connection.getresponse()
        result = response.status, dict(response.getheaders()), response.read()
        connection.close()
        return result

    def test_static_page_serves_local_bundle_and_does_not_expose_files(self) -> None:
        """验证Vue入口使用本地构建资源，并限制文件访问入口。"""
        status, headers, body = self.request("GET", "/", token=False)
        self.assertEqual(status, 200)
        self.assertIn(b'<div id="app"></div>', body)
        self.assertIn(b'src="/app.js"', body)
        self.assertIn(b'href="/app.css"', body)
        for asset in ("/app.js", "/app.css", "/favicon.svg"):
            with self.subTest(asset=asset):
                self.assertEqual(self.request("GET", asset, token=False)[0], 200)
        self.assertIn("frame-ancestors 'none'", headers["Content-Security-Policy"])
        self.assertEqual(self.request("GET", "/.env")[0], 404)
        self.assertEqual(self.request("GET", "/app.mjs")[0], 404)

    def test_token_host_and_origin_are_enforced(self):
        """验证令牌、主机与来源限制。"""
        self.assertEqual(self.request("GET", "/api/session", token=False)[0], 403)
        self.assertEqual(self.request("GET", "/api/session", headers={"X-ASR-Token": "é"})[0], 403)
        self.assertEqual(self.request("GET", "/api/session", headers={"Host": "evil.example"})[0], 403)
        self.assertEqual(self.request("POST", "/api/validate",
                                      headers={"Origin": "https://evil.example"})[0], 403)

    def test_second_server_cannot_reuse_live_port(self):
        """验证已占用端口拒绝第二个服务绑定。"""
        with self.assertRaises(OSError):
            create_server(self.runtime, self.server.server_port)

    def test_cookie_restores_ready_preview_without_creating_a_job(self):
        """验证无令牌URL通过会话Cookie恢复当前预览。"""
        status, headers, _ = self.request("GET", "/", token=False)
        self.assertEqual(status, 200)
        cookie = headers["Set-Cookie"]
        self.assertIn("HttpOnly", cookie)
        self.assertIn("SameSite=Strict", cookie)
        preview = json.loads(self.request("POST", "/api/validate", self.payload)[2])
        self.request("POST", "/api/preview-ready", {"validation_id": preview["validation_id"]})
        status, _, body = self.request("GET", "/api/session", token=False, headers={"Cookie": cookie.split(";", 1)[0]})
        self.assertEqual(status, 200)
        description = json.loads(body)
        self.assertEqual(description["phase"], "preview")
        self.assertEqual(description["preview"]["validation_id"], preview["validation_id"])
        self.assertFalse(self.runtime.path(".state/jobs").exists())

    def test_confirm_returns_authorized_receipt_without_private_inputs(self):
        """验证代码交接返回授权任务，并且回执与输出不包含用户内容。"""
        secret = "synthetic-secret-key"
        context = "只用于测试的内部会议术语"
        self.runtime.path(".env").write_text("DASHSCOPE_API_KEY=" + secret, encoding="utf-8")
        payload = {**self.payload, "auth_mode": "api_key", "enhancement_mode": "context", "context": context}
        preview = json.loads(self.request("POST", "/api/validate", payload)[2])
        self.request("POST", "/api/preview-ready", {"validation_id": preview["validation_id"]})
        with patch("builtins.print") as printed:
            status, _, body = self.request("POST", "/api/confirm", {})
        self.assertEqual(status, 200)
        receipt = json.loads(body)
        self.assertEqual(receipt["auth_mode"], "api_key")
        self.assertEqual(receipt["session_id"], self.server.session.session_id)
        self.assertFalse(receipt["execution_started"])
        self.assertNotIn(secret, body.decode())
        self.assertNotIn(context, body.decode())
        printed.assert_not_called()
        self.assertTrue(job_status(self.runtime, receipt["job_id"])["execution_authorized"])

    def test_browser_cookie_cannot_confirm_or_cancel(self):
        """验证浏览器预览与本机代码交接使用各自的接口权限。"""
        _, headers, _ = self.request("GET", "/", token=False)
        cookie = headers["Set-Cookie"].split(";", 1)[0]
        for endpoint in ("/api/confirm", "/api/cancel"):
            with self.subTest(endpoint=endpoint):
                self.assertEqual(self.request("POST", endpoint, {}, token=False, headers={"Cookie": cookie})[0], 403)
        self.assertFalse(self.runtime.path(".state/jobs").exists())

    def test_direct_upload_and_template(self):
        """验证直接上传与热词模板下载。"""
        status, _, body = self.request("POST", "/api/upload-audio", body=self.audio.read_bytes(),
                                      headers={"Content-Type": "application/octet-stream",
                                               "X-File-Name": quote(self.audio.name)})
        self.assertEqual(status, 200)
        uploaded = json.loads(body)
        self.assertEqual(uploaded["name"], self.audio.name)
        self.assertEqual(self.request("POST", "/api/validate",
                                      {**self.payload, "audio_upload_id": uploaded["upload_id"]})[0], 200)
        self.assertEqual(self.request("GET", "/api/files?directory=data/audio")[0], 404)
        status, _, body = self.request("GET", "/api/hotwords-template")
        self.assertEqual(status, 200)
        workbook = load_workbook(io.BytesIO(body))
        self.assertEqual(workbook["热词"]["A1"].value, "text")
        workbook.close()

    def test_upload_filename_cannot_supply_a_path(self):
        """验证路径形式的上传文件名被拒绝。"""
        status, _, body = self.request("POST", "/api/upload-audio",
                                   headers={"Content-Type": "application/octet-stream",
                                            "X-File-Name": quote("../escape.wav")})
        self.assertEqual(status, 422)
        self.assertIn("所选文件名或格式不符合要求", json.loads(body)["error"])
        self.assertFalse(self.runtime.path("escape.wav").exists())

    def test_missing_api_key_returns_empty_editable_value(self):
        """验证未配置Key时页面可直接填写凭据。"""
        status, _, body = self.request("POST", "/api/api-key", {})
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body), {"ok": True, "value": ""})
        self.assertEqual(self.request("POST", "/api/auth-status", {"auth_mode": "api_key"})[0], 404)

    def test_api_key_save_requires_local_origin_and_session_token(self):
        """验证服务在读取Key正文前拒绝无令牌或跨来源请求。"""
        for options in ({"token": False}, {"headers": {"Origin": "https://evil.example"}}):
            with self.subTest(options=options):
                # 声明正文但先等拒绝回执，避免向已关闭的连接继续发正文。
                headers = {"Content-Type": "application/json",
                           "Content-Length": str(len(b'{"value":"synthetic-secret"}')),
                           **options.get("headers", {})}
                status, _, _ = self.request("POST", "/api/save-api-key",
                                            token=options.get("token", True), headers=headers)
                self.assertEqual(status, 403)
                self.assertFalse(self.runtime.path(".env").exists())

    def test_api_key_save_uses_fixed_workspace_path_without_response_echo(self):
        """验证Key保存固定在工作区.env且回执、预览、配置不包含密钥。"""
        secret = "new-synthetic-secret"
        with patch("builtins.print") as printed:
            status, _, body = self.request("POST", "/api/save-api-key", {
                "value": secret, "path": str(self.runtime.skill_root / ".env"),
            })
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body), {"ok": True})
        self.assertFalse(printed.called)
        self.assertFalse((self.runtime.skill_root / ".env").exists())
        self.assertFalse((self.runtime.workspace / ".env").exists())
        self.assertIn(secret, self.runtime.path(".env").read_text(encoding="utf-8"))
        preview = json.loads(self.request("POST", "/api/validate", {**self.payload, "auth_mode": "api_key"})[2])
        self.request("POST", "/api/preview-ready", {"validation_id": preview["validation_id"]})
        receipt = json.loads(self.request("POST", "/api/confirm", {})[2])
        self.assertNotIn(secret, json.dumps(preview))
        self.assertNotIn(secret, Path(receipt["config_path"]).read_text(encoding="utf-8"))

    def test_api_key_save_rejects_invalid_input_without_overwriting_existing_key(self):
        """验证Key输入错误定位到表单，原有凭据保持有效。"""
        self.runtime.path(".env").write_text("DASHSCOPE_API_KEY=previous-synthetic\n", encoding="utf-8")
        status, _, body = self.request("POST", "/api/save-api-key", {"value": "invalid synthetic"})
        self.assertEqual(status, 422)
        self.assertEqual(json.loads(body)["field"], "auth_mode")
        self.assertNotIn(b"invalid synthetic", body)
        self.assertIn("previous-synthetic", self.runtime.path(".env").read_text(encoding="utf-8"))

    def test_preview_must_return_to_edit_before_changing_api_key(self):
        """验证已就绪预览通过后端返回填写后才能修改凭据。"""
        preview = json.loads(self.request("POST", "/api/validate", self.payload)[2])
        self.request("POST", "/api/preview-ready", {"validation_id": preview["validation_id"]})
        self.assertEqual(self.request("POST", "/api/save-api-key", {"value": "synthetic-key"})[0], 422)
        self.assertFalse(self.runtime.path(".env").exists())
        self.assertEqual(self.request("POST", "/api/edit", {"validation_id": preview["validation_id"]})[0], 200)
        self.assertEqual(self.request("POST", "/api/save-api-key", {"value": "synthetic-key"})[0], 200)

    def test_request_language_applies_to_errors_and_resets_for_chinese(self) -> None:
        """验证请求语言覆盖协议与字段错误，并保持默认中文。"""
        english = {"Accept-Language": "en-US"}
        status, _, body = self.request("GET", "/api/session", token=False, headers=english)
        self.assertEqual(status, 403)
        self.assertIn("Reopen the local page from Codex", json.loads(body)["error"])
        invalid = {**self.payload, "audio_upload_id": "missing"}
        status, _, body = self.request("POST", "/api/validate", invalid, headers=english)
        self.assertEqual(status, 422)
        self.assertEqual(json.loads(body)["error"], "Choose an audio file.")
        for headers in (None, {"Accept-Language": "fr"}):
            with self.subTest(headers=headers):
                status, _, body = self.request("POST", "/api/validate", invalid, headers=headers)
                self.assertEqual(json.loads(body)["error"], "请选择音频文件。")

    def test_concurrent_http_languages_are_independent(self) -> None:
        """验证并发HTTP请求使用各自的提示语言。"""
        invalid = {**self.payload, "audio_upload_id": "missing"}

        def error_in(language: str) -> str:
            """读取指定页面语言下的字段错误。"""
            body = self.request("POST", "/api/validate", invalid,
                                headers={"Accept-Language": language})[2]
            return str(json.loads(body)["error"])

        with ThreadPoolExecutor(max_workers=2) as pool:
            english = pool.submit(error_in, "en")
            chinese = pool.submit(error_in, "zh-CN")
            self.assertEqual(english.result(timeout=5), "Choose an audio file.")
            self.assertEqual(chinese.result(timeout=5), "请选择音频文件。")

    def test_english_context_error_preserves_dynamic_limits(self) -> None:
        """验证上下文超限提示先翻译模板再填入实际数量。"""
        payload = {**self.payload, "enhancement_mode": "context", "context": "词" * 401}
        status, _, body = self.request("POST", "/api/validate", payload,
                                      headers={"Accept-Language": "en"})
        self.assertEqual(status, 422)
        error = json.loads(body)
        self.assertEqual(error["field"], "context")
        self.assertIn("401 characters", error["error"])
        self.assertIn("400-character limit by 1", error["error"])
        self.assertNotIn("词", error["error"])

    def test_english_preview_translates_audio_warnings_and_keeps_filename(self) -> None:
        """验证音频警告随页面语言显示，保留用户文件名。"""
        info = AudioInfo(2, 16000, 7201, self.audio.stat().st_size, "wav", 2)
        with patch("asr_runtime.application.inputs.probe_audio", return_value=info):
            status, _, body = self.request("POST", "/api/validate", self.payload,
                                          headers={"Accept-Language": "en"})
        self.assertEqual(status, 200)
        summary = json.loads(body)["summary"]
        self.assertEqual(summary["audio"]["name"], self.audio.name)
        self.assertEqual(len(summary["warnings"]), 3)
        self.assertIn("mono FLAC copy", summary["warnings"][0])
        self.assertIn("longer than 2 hours", summary["warnings"][1])
        self.assertIn("2 audio tracks", summary["warnings"][2])

    def test_english_hotword_row_error_identifies_template_header(self) -> None:
        """验证工作簿和行级错误都使用英文并保留行号。"""
        workbook = Workbook()
        workbook.active.append(["wrong", "header"])
        content = io.BytesIO()
        workbook.save(content)
        workbook.close()
        status, _, body = self.request("POST", "/api/upload-hotwords", body=content.getvalue(),
                              headers={"Content-Type": "application/octet-stream",
                                       "X-File-Name": quote("热词.xlsx"), "Accept-Language": "en"})
        self.assertEqual(status, 422)
        error = json.loads(body)
        self.assertEqual(error["field"], "hotword_rows")
        self.assertIn("The first row must contain text and weight", error["error"])
        self.assertEqual(error["details"][0], {
            "row": 1, "field": "header", "message": "Use the column headers from the template.",
        })

    def test_hotword_import_returns_invalid_rows_for_inline_correction(self):
        """验证HTTP导入保留错误行，修改后的表格使用相同规则通过检查。"""
        workbook = Workbook()
        for row in (["text", "weight"], ["术语", 4], ["另一个词", 99], ["术语", 5]):
            workbook.active.append(row)
        content = io.BytesIO()
        workbook.save(content)
        workbook.close()
        status, _, body = self.request("POST", "/api/upload-hotwords", body=content.getvalue(),
                                      headers={"Content-Type": "application/octet-stream", "X-File-Name": "words.xlsx"})
        self.assertEqual(status, 200)
        imported = json.loads(body)
        self.assertEqual([(item["row"], item["field"]) for item in imported["issues"]],
                         [(1, "text"), (2, "weight"), (3, "text")])
        self.assertEqual(imported["rows"][1]["weight"], 99)
        self.assertFalse(list(self.server.session.upload_directory.glob("*.xlsx*")))
        imported["rows"][1]["weight"] = "4"
        imported["rows"].pop(2)
        status, _, body = self.request("POST", "/api/validate-hotwords", {"rows": imported["rows"]})
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body)["issues"], [])
        self.assertEqual(json.loads(body)["count"], 2)

    def test_hotword_request_accepts_model_limit_beyond_old_body_limit(self):
        """验证2000条词表可通过HTTP检查，避免被旧32KiB请求上限误拒绝。"""
        rows = [{"text": f"术语{number}", "weight": "4"} for number in range(2000)]
        self.assertGreater(len(json.dumps({"rows": rows}).encode()), 32 * 1024)
        status, _, body = self.request("POST", "/api/validate-hotwords", {"rows": rows})
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body)["count"], 2000)
        self.assertEqual(json.loads(body)["issues"], [])

    def test_return_to_edit_invalidates_preview_without_creating_a_job(self):
        """验证返回修改使旧预览失效，重复填写始终没有执行任务。"""
        preview = json.loads(self.request("POST", "/api/validate", self.payload)[2])
        self.request("POST", "/api/preview-ready", {"validation_id": preview["validation_id"]})
        status, _, body = self.request("POST", "/api/edit", {"validation_id": preview["validation_id"]})
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body)["configuration"], self.payload)
        self.assertEqual(self.request("POST", "/api/preview-ready", {"validation_id": preview["validation_id"]})[0], 422)
        self.assertEqual(self.request("POST", "/api/confirm", {})[0], 422)
        self.assertFalse(self.runtime.path(".state/jobs").exists())
        revised = json.loads(self.request("POST", "/api/validate", self.payload)[2])
        self.assertNotEqual(revised["validation_id"], preview["validation_id"])
        self.assertFalse(self.runtime.path(".state/jobs").exists())
        self.assertEqual(self.request("POST", "/api/reopen", {})[0], 404)

    def test_english_api_key_errors_guide_page_input_without_echoing_key(self) -> None:
        """验证Key空值可填写，保存校验英文提示保持凭据脱敏。"""
        english = {"Accept-Language": "en"}
        status, _, body = self.request("POST", "/api/api-key", {}, headers=english)
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body), {"ok": True, "value": ""})
        status, _, body = self.request("POST", "/api/save-api-key", {"value": ""}, headers=english)
        self.assertEqual(status, 422)
        self.assertEqual(json.loads(body)["error"], "Enter your API key.")
        secret = "synthetic secret-key"
        status, _, body = self.request("POST", "/api/save-api-key", {"value": secret}, headers=english)
        self.assertEqual(status, 422)
        self.assertIn("spaces or line breaks", json.loads(body)["error"])
        self.assertNotIn(secret, body.decode("utf-8"))

    def test_english_template_keeps_schema_and_translates_example(self) -> None:
        """验证英文模板沿用现有热词结构并使用英文示例。"""
        status, _, body = self.request("GET", "/api/hotwords-template", headers={"Accept-Language": "en"})
        self.assertEqual(status, 200)
        workbook = load_workbook(io.BytesIO(body))
        try:
            sheet = workbook["热词"]
            self.assertEqual([sheet["A1"].value, sheet["B1"].value], ["text", "weight"])
            self.assertEqual(sheet["A2"].value, "Example term")
        finally:
            workbook.close()
