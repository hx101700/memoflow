"""验证固定目录导出、同名替换和格式失败处理。"""

import hashlib
import json
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

from asr_runtime import MODEL
from asr_runtime.utils.job_files import read_delivery
from asr_runtime.utils.documents import DocumentError, write_xlsx
from asr_runtime.utils.files import write_json_atomic
from asr_runtime.utils.environment import SetupError
from asr_runtime.utils.results import load_transcript
from asr_runtime.utils.documents import timestamp
from asr_runtime.application.transcription import export_job, job_status, transcribe
from tests.support import RuntimeTestCase


class DeliveryTests(RuntimeTestCase):
    def setUp(self):
        """准备已保存任务、合成JSON和文档目录。"""
        super().setUp()
        self.job_id = "a" * 32
        self.root = self.runtime.path(f".state/jobs/{self.job_id}")
        (self.root / "execution").mkdir(parents=True)
        base = self.runtime.output_root / self.job_id
        self.config = {"schema_version": 1, "job_id": self.job_id, "model": MODEL,
                       "region": "cn-beijing", "status": "CONFIGURED",
                       "execution_authorized": True, "confirmed_at": "synthetic-time",
                       "json_directory": str(base / "json"),
                       "document_directory": str(base / "documents"),
                       "audio": {"name": "合成录音.wav"}}
        self.save_config()
        self.json_path = base / "json/transcription.json"
        self.json_path.parent.mkdir(parents=True)
        self.json_path.write_text(json.dumps({
            "file_url": "https://example.invalid/private-result?signature=synthetic-secret",
            "transcripts": [{"channel_id": 0, "sentences": [
                {"text": "合成文本。", "begin_time": 1, "end_time": 1501, "speaker_id": 0},
                {"text": "", "begin_time": 1501, "end_time": 1501},
            ]}, {"channel_id": 2, "sentences": [
                {"text": "=HYPERLINK(\"https://example.invalid\")", "begin_time": 0, "end_time": 1234},
            ]}]
        }, ensure_ascii=False), encoding="utf-8")
        self.save_execution()
        cloud = patch("asr_runtime.application.transcription.run_recognition", side_effect=AssertionError("local export must not call BL"))
        self.cloud = cloud.start()
        self.addCleanup(cloud.stop)

    def save_config(self):
        """保存合成任务配置及对应内容摘要。"""
        content = json.dumps(self.config, ensure_ascii=False).encode("utf-8")
        (self.root / "config.json").write_bytes(content)
        (self.root / "config.sha256").write_text(hashlib.sha256(content).hexdigest(), encoding="ascii")

    def save_execution(self, **changes):
        """保存测试用执行状态与结果摘要。"""
        report = {"job_id": self.job_id, "status": "JSON_READY", "documents_ready": False,
                  "cloud_outcome": "result_received", "json_path": str(self.json_path),
                  "result": load_transcript(self.json_path).summary()}
        report.update(changes)
        (self.root / "execution/status.json").write_text(json.dumps(report), encoding="utf-8")

    def test_existing_json_exports_without_key_audio_or_bl_and_preserves_source(self):
        """验证已有JSON独立导出并保留原结果。"""
        original = self.json_path.read_bytes()
        report = export_job(self.runtime, self.job_id)
        self.assertTrue(report["documents_ready"])
        self.assertEqual(report["delivery"]["status"], "COMPLETE")
        self.assertEqual(set(report["delivery"]["files"]), {"xlsx", "docx", "md"})
        for record in report["delivery"]["files"].values():
            content = Path(record["path"]).read_bytes()
            self.assertEqual(len(content), record["bytes"])
            self.assertNotIn("sha256", record)
        self.assertEqual(report["result"]["sha256"], hashlib.sha256(original).hexdigest())
        self.assertEqual(self.json_path.read_bytes(), original)
        self.cloud.assert_not_called()
        self.assertEqual(job_status(self.runtime, self.job_id), report)

    def test_one_format_failure_retains_other_formats_and_never_retries(self):
        """验证单格式失败保留其他成品且各格式调用一次。"""
        with patch("asr_runtime.application.delivery.write_docx", side_effect=RuntimeError("private transcript in exception")) as writer:
            report = export_job(self.runtime, self.job_id)
        writer.assert_called_once()
        self.assertFalse(report["documents_ready"])
        self.assertEqual(report["status"], "JSON_READY")
        self.assertEqual(report["cloud_outcome"], "result_received")
        self.assertEqual(report["delivery"]["status"], "PARTIAL")
        self.assertEqual(report["delivery"]["files"]["docx"]["error_type"], "RuntimeError")
        self.assertIn("程序发生异常", report["delivery"]["files"]["docx"]["message"])
        self.assertNotIn("磁盘", report["delivery"]["files"]["docx"]["message"])
        self.assertNotIn("private transcript", json.dumps(report))
        for extension in ("xlsx", "md"):
            self.assertTrue(Path(report["delivery"]["files"][extension]["path"]).is_file())
        self.assertEqual(transcribe(self.runtime, self.job_id), report)
        self.cloud.assert_not_called()

    def test_known_document_error_and_filesystem_error_have_specific_messages(self):
        """验证格式错误和文件系统错误有对应说明。"""
        for error, message in ((DocumentError("已知格式限制"), "已知格式限制"),
                               (OSError("private path"), "目录权限")):
            with self.subTest(error=type(error).__name__), \
                 patch("asr_runtime.application.delivery.write_docx", side_effect=error):
                report = export_job(self.runtime, self.job_id)
            failure = report["delivery"]["files"]["docx"]
            self.assertIn(message, failure["message"])
            self.assertNotIn("private path", json.dumps(report))
            self.assertEqual(report["delivery"]["status"], "PARTIAL")

    def test_unexpected_serializer_error_is_not_reported_as_disk_failure(self):
        """验证序列化异常返回对应的程序错误说明。"""
        with patch("asr_runtime.utils.documents.Workbook.save", side_effect=ValueError("private transcript")):
            report = export_job(self.runtime, self.job_id)
        failure = report["delivery"]["files"]["xlsx"]
        self.assertEqual(failure["error_type"], "ValueError")
        self.assertIn("程序发生异常", failure["message"])
        self.assertNotIn("目录", failure["message"])
        self.assertNotIn("private transcript", json.dumps(report))
        self.assertEqual(report["delivery"]["status"], "PARTIAL")

    def test_explicit_reexport_replaces_the_same_three_files(self):
        """验证重导使用固定目录并替换三个同名文件。"""
        first = export_job(self.runtime, self.job_id)["delivery"]
        expected = {Path(self.config["document_directory"]) / f"transcription.{ext}"
                    for ext in ("xlsx", "docx", "md")}
        for record in first["files"].values():
            Path(record["path"]).write_bytes(b"synthetic edited document")
        second = export_job(self.runtime, self.job_id)["delivery"]
        self.assertEqual(second["status"], "COMPLETE")
        for extension, record in first["files"].items():
            self.assertEqual(record["path"], second["files"][extension]["path"])
            self.assertNotEqual(Path(record["path"]).read_bytes(), b"synthetic edited document")
        self.assertEqual(set(Path(self.config["document_directory"]).iterdir()), expected)
        self.assertTrue((self.root / "delivery/status.json").is_file())
        self.assertFalse((self.root / "exports").exists())
        self.assertEqual(read_delivery(self.root), second)
        self.cloud.assert_not_called()

    def test_overlapping_exports_do_not_replace_each_others_temporary_files(self):
        """验证两次导出交错时各自发布成功，不互删临时文档。"""
        ready, release = threading.Event(), threading.Event()
        lock = threading.Lock()
        paths = []

        def pause_first_xlsx(transcript, path, **metadata):
            """首轮真实Excel生成后暂停，等待第二轮完成再发布。"""
            with lock:
                paths.append(path)
                first = len(paths) == 1
            write_xlsx(transcript, path, **metadata)
            if first:
                ready.set()
                self.assertTrue(release.wait(5), "测试未释放首轮导出")

        with patch("asr_runtime.application.delivery.write_xlsx", side_effect=pause_first_xlsx), \
                ThreadPoolExecutor(max_workers=2) as pool:
            first = pool.submit(export_job, self.runtime, self.job_id)
            try:
                self.assertTrue(ready.wait(5), "首轮Excel尚未生成")
                second = pool.submit(export_job, self.runtime, self.job_id).result(timeout=5)
            finally:
                release.set()
            first_report = first.result(timeout=5)
        self.assertNotEqual(paths[0], paths[1])
        for report in (first_report, second, job_status(self.runtime, self.job_id)):
            self.assertEqual(report["delivery"]["status"], "COMPLETE")
            self.assertTrue(report["documents_ready"])
        destination = Path(self.config["document_directory"])
        self.assertEqual({path.name for path in destination.iterdir()},
                         {"transcription.xlsx", "transcription.docx", "transcription.md"})
        self.cloud.assert_not_called()

    def test_overlapping_records_keep_their_own_temporary_files(self):
        """验证同一JSON目标的交错发布不删除另一调用的临时记录。"""
        destination = self.root / "audit-record.json"
        ready, release = threading.Event(), threading.Event()
        lock = threading.Lock()
        paths = []
        replace = Path.replace

        def pause_first_replace(path, target):
            """暂停首轮记录替换，让第二轮先完成发布。"""
            with lock:
                paths.append(path)
                first = len(paths) == 1
            if first:
                ready.set()
                self.assertTrue(release.wait(5), "测试未释放首轮记录")
            return replace(path, target)

        with patch.object(Path, "replace", autospec=True, side_effect=pause_first_replace), \
                ThreadPoolExecutor(max_workers=2) as pool:
            first = pool.submit(write_json_atomic, destination, {"writer": "first"})
            try:
                self.assertTrue(ready.wait(5), "首轮记录尚未写入")
                second = pool.submit(write_json_atomic, destination, {"writer": "second"})
                second.result(timeout=5)
            finally:
                release.set()
            first.result(timeout=5)
        self.assertNotEqual(paths[0], paths[1])
        self.assertEqual(json.loads(destination.read_text(encoding="utf-8")), {"writer": "first"})
        self.assertFalse(list(self.root.glob("audit-record.json.*.tmp")))

    def test_failed_record_replacement_preserves_target_and_cleans_its_temporary_file(self):
        """验证记录替换失败保留原内容并移除本次临时记录。"""
        destination = self.root / "audit-record.json"
        destination.write_text('{"writer":"original"}', encoding="utf-8")
        with patch.object(Path, "replace", side_effect=PermissionError("synthetic target locked")), \
                self.assertRaises(PermissionError):
            write_json_atomic(destination, {"writer": "new"})
        self.assertEqual(json.loads(destination.read_text(encoding="utf-8")), {"writer": "original"})
        self.assertFalse(list(self.root.glob("audit-record.json.*.tmp")))

    def test_failed_reexport_preserves_the_existing_document(self):
        """验证重导写入失败时保留已有成品并完成其他格式。"""
        first = export_job(self.runtime, self.job_id)["delivery"]
        document = Path(first["files"]["docx"]["path"])
        original = document.read_bytes()

        def fail_docx(transcript, path, **metadata):
            """模拟临时Word写到一半时发生文件错误。"""
            path.write_bytes(b"synthetic incomplete document")
            raise OSError("synthetic write failure")

        with patch("asr_runtime.application.delivery.write_docx", side_effect=fail_docx) as writer:
            report = export_job(self.runtime, self.job_id)
        writer.assert_called_once()
        self.assertEqual(document.read_bytes(), original)
        self.assertEqual(report["delivery"]["status"], "PARTIAL")
        self.assertEqual(report["delivery"]["files"]["docx"]["status"], "FAILED")
        self.assertFalse(list(document.parent.glob("*.partial-*")))
        for extension in ("xlsx", "md"):
            self.assertEqual(report["delivery"]["files"][extension]["status"], "READY")
        self.assertEqual(job_status(self.runtime, self.job_id), report)
        self.cloud.assert_not_called()

    def test_locked_target_preserves_the_existing_document(self):
        """验证同名文件被占用时报告替换失败并保留旧文件。"""
        first = export_job(self.runtime, self.job_id)["delivery"]
        document = Path(first["files"]["docx"]["path"])
        original = document.read_bytes()
        replace = Path.replace

        def refuse_word_replacement(source, target):
            """模拟目标Word被其他程序占用，其余文件照常发布。"""
            if Path(target) == document:
                raise PermissionError("synthetic file in use")
            return replace(source, target)

        with patch.object(Path, "replace", autospec=True, side_effect=refuse_word_replacement):
            report = export_job(self.runtime, self.job_id)
        self.assertEqual(document.read_bytes(), original)
        self.assertEqual(report["delivery"]["status"], "PARTIAL")
        self.assertEqual(report["delivery"]["files"]["docx"]["error_type"], "PermissionError")
        self.cloud.assert_not_called()

    def test_changed_json_after_recognition_is_rejected(self):
        """验证识别后JSON内容改变时拒绝导出。"""
        self.json_path.write_bytes(self.json_path.read_bytes() + b"\n")
        with self.assertRaisesRegex(SetupError, "发生变化"):
            export_job(self.runtime, self.job_id)
        self.assertFalse((self.root / "delivery").exists())

    def test_result_without_digest_is_not_exported_or_resubmitted(self):
        """验证结果摘要缺失时返回导出来源错误。"""
        self.save_execution(result={"audio_tracks": 2, "sentences": 3})
        with self.assertRaisesRegex(SetupError, "缺少结果JSON摘要"):
            export_job(self.runtime, self.job_id)
        self.assertFalse((self.root / "delivery").exists())
        self.cloud.assert_not_called()

    def test_no_export_when_recognition_did_not_finish(self):
        """验证识别进行中拒绝导出。"""
        self.save_execution(status="RUNNING")
        with self.assertRaisesRegex(SetupError, "尚无"):
            export_job(self.runtime, self.job_id)
        self.assertFalse((self.root / "delivery").exists())

    def test_unwritable_output_is_local_failure_and_preserves_json(self):
        """验证保存目录写入失败时保留JSON并报告导出失败。"""
        Path(self.config["document_directory"]).write_text("existing file", encoding="utf-8")
        report = export_job(self.runtime, self.job_id)
        self.assertEqual(report["delivery"]["status"], "FAILED")
        self.assertEqual(report["status"], "JSON_READY")
        self.assertFalse(report["documents_ready"])
        self.assertTrue(self.json_path.is_file())

    def test_export_destination_in_skill_is_rejected_and_json_is_preserved(self):
        """验证最终文档目录在Skill中时保留JSON并报告本地失败。"""
        destination = self.runtime.skill_root / self.job_id / "documents"
        self.config["document_directory"] = str(destination)
        self.save_config()
        before = self.json_path.read_bytes()
        report = export_job(self.runtime, self.job_id)
        self.assertEqual(report["status"], "JSON_READY")
        self.assertEqual(report["delivery"]["status"], "FAILED")
        self.assertEqual(report["delivery"]["error_type"], "SetupError")
        self.assertIn("Skill安装目录", report["delivery"]["message"])
        self.assertFalse(report["documents_ready"])
        self.assertFalse(destination.parent.exists())
        self.assertEqual(self.json_path.read_bytes(), before)
        self.cloud.assert_not_called()

    def test_saved_result_exports_with_its_recorded_model(self):
        """验证更新后的本地重导保留任务原模型，不改配置或重新识别。"""
        recorded_model = "qwen-audio-3.0-asr-flash-filetrans"
        self.config["model"] = recorded_model
        self.save_config()
        original_config = (self.root / "config.json").read_bytes()
        original_checksum = (self.root / "config.sha256").read_bytes()
        original_json = self.json_path.read_bytes()
        report = export_job(self.runtime, self.job_id)
        self.assertTrue(report["documents_ready"])
        self.assertEqual(job_status(self.runtime, self.job_id), report)
        files = report["delivery"]["files"]
        with Path(files["xlsx"]["path"]).open("rb") as stream:
            workbook = load_workbook(stream, read_only=True)
            try:
                self.assertEqual(workbook["转写明细"]["A2"].value,
                                 f"模型：{recorded_model}    任务：{self.job_id}")
            finally:
                workbook.close()
        document = Document(files["docx"]["path"])
        self.assertEqual(document.paragraphs[1].text, f"模型：{recorded_model}")
        markdown_model = recorded_model.replace("-", "\\-").replace(".", "\\.")
        self.assertIn(f"模型：{markdown_model}",
                      Path(files["md"]["path"]).read_text(encoding="utf-8"))
        self.assertEqual((self.root / "config.json").read_bytes(), original_config)
        self.assertEqual((self.root / "config.sha256").read_bytes(), original_checksum)
        self.assertEqual(self.json_path.read_bytes(), original_json)
        self.cloud.assert_not_called()

    def test_unreadable_delivery_record_reports_unknown(self):
        """验证导出记录损坏时返回结果未知状态。"""
        export_job(self.runtime, self.job_id)
        (self.root / "delivery/status.json").write_text("{", encoding="utf-8")
        report = job_status(self.runtime, self.job_id)
        self.assertEqual(report["delivery"]["status"], "OUTCOME_UNKNOWN")
        self.assertFalse(report["documents_ready"])

    def test_delivery_files_list_is_reported_as_unknown(self) -> None:
        """验证成品映射损坏为列表时保留JSON并返回未知交付状态。"""
        directory = self.root / "delivery"
        directory.mkdir()
        (directory / "status.json").write_text(json.dumps({
            "job_id": self.job_id, "status": "COMPLETE", "message": "synthetic",
            "files": ["xlsx", "docx", "md"],
        }), encoding="utf-8")
        report = job_status(self.runtime, self.job_id)
        self.assertEqual(report["status"], "JSON_READY")
        self.assertEqual(report["delivery"]["status"], "OUTCOME_UNKNOWN")
        self.assertFalse(report["documents_ready"])
        self.assertTrue(self.json_path.is_file())
        self.cloud.assert_not_called()

    def test_status_write_failure_still_reports_saved_json_without_reidentification(self):
        """验证状态写入失败仍说明JSON已保存。"""
        with patch("asr_runtime.application.delivery.save_record", side_effect=OSError("synthetic disk full")):
            report = export_job(self.runtime, self.job_id)
        self.assertEqual(report["status"], "JSON_READY")
        self.assertEqual(report["cloud_outcome"], "result_received")
        self.assertEqual(report["delivery"]["status"], "OUTCOME_UNKNOWN")
        self.assertFalse(report["documents_ready"])
        self.assertTrue(Path(report["json_path"]).is_file())
        self.cloud.assert_not_called()

    def test_export_returns_current_delivery_without_status_lookup(self):
        """验证导出直接返回当前执行结果。"""
        expected = {"status": "PARTIAL", "message": "synthetic"}
        with patch("asr_runtime.application.transcription.export_documents", return_value=expected), \
             patch("asr_runtime.application.transcription.read_delivery", side_effect=AssertionError("导出不重新读取状态")):
            report = export_job(self.runtime, self.job_id)
        self.assertEqual(report["delivery"], expected)

    def test_loader_preserves_order_empty_sentence_ids_and_precise_milliseconds(self):
        """验证JSON解析保留原顺序、空句和毫秒时间。"""
        transcript = load_transcript(self.json_path)
        self.assertEqual([s.begin_ms for s in transcript.sentences], [1, 1501, 0])
        self.assertEqual([s.channel_id for s in transcript.sentences], [0, 0, 2])
        self.assertEqual([s.track_index for s in transcript.sentences], [1, 1, 2])
        self.assertEqual(transcript.sentences[1].text, "")
        self.assertIsNone(transcript.sentences[1].speaker_id)
        self.assertEqual(timestamp(25 * 3600000 + 1501), "25:00:01.501")

    def test_channel_or_sentence_schema_errors_are_not_guessed(self):
        """验证异常结果字段返回结构检查错误。"""
        for change in ({"channel_id": True, "sentences": []}, {"sentences": [None]},
                       {"sentences": [{"text": "text", "start": 0, "end": 1}]}):
            self.json_path.write_text(json.dumps({"transcripts": [change]}), encoding="utf-8")
            with self.assertRaises(SetupError):
                load_transcript(self.json_path)
