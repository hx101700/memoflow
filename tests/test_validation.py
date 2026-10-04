import hashlib
import os
import subprocess
import wave
import zipfile
from dataclasses import replace
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from asr_runtime.models import AudioInfo
from asr_runtime.application.rules import (
    AUDIO_SUFFIXES, MAX_UPLOAD_BYTES, ValidationError, validate_context, validate_hotword_rows,
)
from asr_runtime.utils.files import FileError, file_fingerprint, resolve_input
from asr_runtime.utils.i18n import language_scope
from asr_runtime.application.inputs import import_hotwords, validate_audio
from tests.support import RuntimeTestCase


class ValidationTests(RuntimeTestCase):
    def setUp(self):
        """准备媒体样本目录和输入校验环境。"""
        super().setUp()
        self.data = self.runtime.path("data")
        self.data.mkdir()

    def audio(self, channels=1):
        """生成指定声道数的合成WAV样本。"""
        path = self.data / "合成.wav"
        with wave.open(str(path), "wb") as output:
            output.setnchannels(channels)
            output.setsampwidth(2)
            output.setframerate(8000)
            output.writeframes(b"\x00\x00" * 8000 * channels)
        return path

    def hotwords(self, rows, *, headers=("text", "weight"), sheet_name="热词"):
        """生成指定词条、表头和工作表的热词Excel。"""
        path = self.data / "合成热词.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def test_resolve_accepts_absolute_local_copy(self):
        """验证接受输入根内的绝对文件路径。"""
        path = self.audio()
        self.assertEqual(resolve_input(self.runtime.root, path, AUDIO_SUFFIXES), path)

    def test_resolve_rejects_relative_outside_missing_and_wrong_kind(self):
        """验证非法路径或文件类型返回输入错误。"""
        self.audio()
        unsupported = self.data / "unsupported.txt"
        unsupported.write_bytes(b"synthetic")
        for value in ("", "data/合成.wav", self.runtime.root.parent / "outside.wav",
                      self.data, self.data / "missing.wav", unsupported):
            with self.subTest(value=value), self.assertRaises(FileError):
                resolve_input(self.runtime.root, value, AUDIO_SUFFIXES)

    def test_resolve_rejects_junction_or_symlink_outside_input_root(self):
        """验证指向输入根外的链接或联接被拒绝。"""
        scoped_root = self.runtime.path("scoped")
        scoped_root.mkdir()
        outside = self.data
        self.audio()
        link = scoped_root / "linked"
        if os.name == "nt":
            result = subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(outside)],
                                    capture_output=True, creationflags=subprocess.CREATE_NO_WINDOW)
            self.assertEqual(result.returncode, 0)
        else:
            link.symlink_to(outside, target_is_directory=True)
        with self.assertRaises(FileError):
            resolve_input(scoped_root, link / "合成.wav", AUDIO_SUFFIXES)

    def test_fingerprint_detects_same_size_content_change(self):
        """验证同大小文件内容变化会改变指纹。"""
        path = self.data / "test.txt"
        path.write_bytes(b"abc")
        first = file_fingerprint(path)
        self.assertEqual(first["sha256"], hashlib.sha256(b"abc").hexdigest())
        path.write_bytes(b"def")
        second = file_fingerprint(path)
        self.assertEqual(first["size_bytes"], second["size_bytes"])
        self.assertNotEqual(first["sha256"], second["sha256"])

    def test_audio_mono_uses_original_and_stereo_only_plans_conversion(self):
        """验证单声道引用原文件且立体声预览包含转换计划。"""
        path = self.audio()
        mono = validate_audio(self.runtime.root, path, True)
        self.assertFalse(mono["requires_mono"])
        self.assertEqual(mono["metadata"]["duration_seconds"], 1.0)
        path = self.audio(channels=2)
        before = path.read_bytes()
        stereo = validate_audio(self.runtime.root, path, True)
        self.assertTrue(stereo["requires_mono"])
        self.assertIn("保留原文件", stereo["warnings"][0])
        self.assertFalse(validate_audio(self.runtime.root, path, False)["requires_mono"])
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(list(self.data.iterdir()), [path])

    def test_audio_rejects_empty_corrupt_and_unsupported_extension(self):
        """验证空、损坏或扩展名异常的音频返回输入错误。"""
        for name, contents in (("empty.wav", b""), ("broken.wav", b"not a wave"),
                               ("file.txt", b"not audio")):
            path = self.data / name
            path.write_bytes(contents)
            with self.subTest(name=name), self.assertRaises(ValidationError):
                validate_audio(self.runtime.root, path, True)

    def test_audio_rejects_same_length_modification_during_probe(self):
        """验证探测期间同长度改写音频会被发现。"""
        path = self.audio()
        before = path.stat()

        def changing_probe(source):
            """改写文件并返回媒体信息以模拟探测期间变化。"""
            content = source.read_bytes()
            source.write_bytes(content[:-2] + b"\x01\x01")
            os.utime(source, ns=(before.st_atime_ns, before.st_mtime_ns + 1_000_000))
            return AudioInfo(1, 8000, 1.0, before.st_size, "wav", 1)

        with patch("asr_runtime.application.inputs.probe_audio", side_effect=changing_probe):
            with self.assertRaisesRegex(ValidationError, "校验期间发生变化"):
                validate_audio(self.runtime.root, path, True)

    def test_audio_duration_and_actual_container_boundaries(self):
        """验证时长和实际容器格式的边界规则。"""
        path = self.audio()
        normal = AudioInfo(1, 8000, 1.0, path.stat().st_size, "wav", 1)
        for duration in (None, 0, -1, float("nan"), float("inf"), 43200.001):
            with self.subTest(duration=duration), patch("asr_runtime.application.inputs.probe_audio",
                    return_value=replace(normal, duration_seconds=duration)):
                with self.assertRaises(ValidationError):
                    validate_audio(self.runtime.root, path, True)
        with patch("asr_runtime.application.inputs.probe_audio", return_value=replace(normal, format_name="aiff")):
            with self.assertRaisesRegex(ValidationError, "实际媒体格式"):
                validate_audio(self.runtime.root, path, False)
        with patch("asr_runtime.application.inputs.probe_audio",
                   return_value=replace(normal, duration_seconds=43200, audio_tracks=2)):
            result = validate_audio(self.runtime.root, path, True)
            self.assertEqual(len(result["warnings"]), 2)
            self.assertTrue(any("索引0" in text for text in result["warnings"]))

    def test_upload_size_applies_to_actual_upload_not_source_to_be_merged(self):
        """验证上传大小规则依据原文件或混音后副本。"""
        path = self.audio(channels=2)
        big = AudioInfo(2, 8000, 1.0, MAX_UPLOAD_BYTES + 1, "wav", 1)
        with patch("asr_runtime.application.inputs.probe_audio", return_value=big):
            with self.assertRaisesRegex(ValidationError, "1 GB"):
                validate_audio(self.runtime.root, path, False)
            with patch("asr_runtime.application.inputs.file_fingerprint",
                       return_value={"size_bytes": big.size_bytes, "mtime_ns": path.stat().st_mtime_ns,
                                     "sha256": "synthetic"}):
                self.assertTrue(validate_audio(self.runtime.root, path, True)["requires_mono"])

    def test_context_counts_unicode_characters_and_preserves_text(self):
        """验证上下文按Unicode字符计数且保留原文。"""
        text = "🙂" * 398 + " \n"
        self.assertEqual(validate_context(text), text)
        for invalid in ("", " \n", "测" * 401, "a\x00b", "a\ud800", None):
            with self.subTest(length=len(invalid) if invalid is not None else None):
                with self.assertRaises(ValidationError) as caught:
                    validate_context(invalid)
                self.assertEqual(caught.exception.field, "context")

    def test_hotwords_skips_blank_excel_rows_and_numbers_current_array(self):
        """验证中文表头可用，导入去除空行，行号由当前数组决定。"""
        path = self.hotwords([("语音实验室", 4), (None, None), ("语音实验室", 4), ("hello world", 2)],
                             headers=("热词", "权重"))
        before = path.read_bytes()
        with patch("asr_runtime.application.inputs.file_fingerprint", side_effect=AssertionError("热词导入不计算文件SHA")):
            result = import_hotwords(path)
        self.assertEqual([issue["row"] for issue in result["issues"]], [1, 2])
        self.assertEqual(len(result["rows"]), 3)
        result["rows"].pop(1)
        checked = validate_hotword_rows(result["rows"])
        self.assertEqual(checked["vocabulary"], {"语音实验室": 4, "hello world": 2})
        self.assertEqual(checked["warnings"], [])
        self.assertEqual(len(result["warnings"]), 1)
        self.assertEqual(path.read_bytes(), before)

    def test_hotwords_reports_formula_conflict_empty_word_and_invalid_weights_by_row(self):
        """验证公式、冲突和非法内容按行报告。"""
        path = self.hotwords([
            ("first", 4), ("first", 3), ("=1+1", 4), (None, 2), ("word", True),
            ("second", 2.5), ("third", "4.0"), (" 热词", 3), ("tab\tword", 2),
        ])
        imported = import_hotwords(path)
        details = imported["issues"]
        self.assertEqual({error["row"] for error in details}, set(range(1, 10)))
        self.assertTrue(all(set(error) == {"row", "field", "message"} for error in details))
        with self.assertRaises(ValidationError) as caught:
            validate_hotword_rows(imported["rows"])
        self.assertEqual(caught.exception.field, "hotword_rows")
        self.assertEqual(caught.exception.details, details)

    def test_hotword_length_rules(self):
        """验证中英文热词长度规则。"""
        valid = self.hotwords([("汉" * 15, 1), ("a b c d e f g", 5)])
        self.assertEqual(validate_hotword_rows(import_hotwords(valid)["rows"])["count"], 2)
        invalid = self.hotwords([("汉" * 16, 1), ("a b c d e f g h", 5)])
        self.assertEqual([error["row"] for error in import_hotwords(invalid)["issues"]], [1, 2])

    def test_fixed_model_accepts_super_words_and_limits_their_count(self):
        """验证固定模型支持超级热词并限制其数量。"""
        path = self.hotwords([(f"term{i}", 50) for i in range(50)])
        self.assertEqual(validate_hotword_rows(import_hotwords(path)["rows"])["vocabulary"],
                         {f"term{i}": 50 for i in range(50)})
        path = self.hotwords([(f"term{i}", 50) for i in range(51)])
        self.assertEqual(import_hotwords(path)["issues"][0]["row"], 51)

    def test_hotword_count_limit(self):
        """验证即时热词总数量上限。"""
        path = self.hotwords([(f"term{i}", 4) for i in range(2000)])
        self.assertEqual(validate_hotword_rows(import_hotwords(path)["rows"])["count"], 2000)
        path = self.hotwords([(f"term{i}", 4) for i in range(2001)])
        self.assertEqual(import_hotwords(path)["issues"][0]["row"], 2001)

    def test_hotwords_rejects_invalid_header_extra_columns_and_corrupt_file(self):
        """验证异常表头、多列和损坏词表在导入时被拒绝。"""
        for headers, rows in [(("word", "weight"), [("hello", 4)]),
                              (("text", "weight", "extra"), [("hello", 4, "x")])]:
            with self.subTest(headers=headers, row_count=len(rows)), self.assertRaises(ValidationError):
                import_hotwords(self.hotwords(rows, headers=headers))
        path = self.data / "bad.xlsx"
        path.write_text("not a spreadsheet")
        with self.assertRaises(ValidationError):
            import_hotwords(path)

    def test_hotwords_bounds_archive_size_and_decompressed_size(self):
        """验证词表原始大小与解压大小上限。"""
        path = self.hotwords([("test", 4)])
        with patch("asr_runtime.utils.hotwords.MAX_XLSX_BYTES", 1):
            with self.assertRaisesRegex(ValidationError, "文件上限"):
                import_hotwords(path)
        with patch("asr_runtime.utils.hotwords.MAX_XLSX_UNCOMPRESSED_BYTES", 1):
            with self.assertRaisesRegex(ValidationError, "解压内容"):
                import_hotwords(path)

    def test_hotwords_rejects_xml_entities(self):
        """验证Excel中的XML实体被拒绝。"""
        path = self.hotwords([("test", 4)])
        with zipfile.ZipFile(path) as original:
            entries = {name: original.read(name) for name in original.namelist()}
        entries["xl/worksheets/sheet1.xml"] = (
            b'<!DOCTYPE worksheet [<!ENTITY word "expanded">]>'
            b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            b'<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&word;</t></is></c>'
            b'</row></sheetData></worksheet>'
        )
        with zipfile.ZipFile(path, "w") as modified:
            for name, contents in entries.items():
                modified.writestr(name, contents)
        with self.assertRaises(ValidationError):
            import_hotwords(path)
        path.unlink()  # 解析被拒绝后，不依赖GC才释放Windows文件占用。

    def test_hotwords_rejects_ambiguous_sheets_and_discloses_explicit_selection(self):
        """验证多工作表选择规则并说明选定工作表。"""
        path = self.data / "sheets.xlsx"
        workbook = Workbook()
        workbook.active.title = "one"
        workbook.create_sheet("two")
        workbook.save(path)
        with self.assertRaisesRegex(ValidationError, "多个工作表"):
            import_hotwords(path)
        sheet = workbook.create_sheet("热词")
        sheet.append(["text", "weight"])
        sheet.append(["test", 4])
        workbook.save(path)
        workbook.close()
        result = import_hotwords(path)
        self.assertEqual(validate_hotword_rows(result["rows"])["vocabulary"], {"test": 4})
        self.assertTrue(any("其他工作表" in warning for warning in result["warnings"]))

    def test_hotwords_rejects_chart_sheet_with_field_error_and_releases_file(self) -> None:
        """验证图表工作表返回双语热词字段错误，并立即释放Excel文件。"""
        path = self.data / "chart-sheet.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "data"
        sheet.append(["text", "weight"])
        sheet.append(["test", 4])
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        workbook.create_chartsheet(title="热词").add_chart(chart)
        workbook.save(path)
        workbook.close()

        for language, message in (
            ("zh-CN", "请使用普通工作表填写热词，不支持图表工作表。"),
            ("en", "Enter hotwords in a regular worksheet. Chart sheets are not supported."),
        ):
            with self.subTest(language=language), language_scope(language):
                with self.assertRaises(ValidationError) as caught:
                    import_hotwords(path)
                self.assertEqual(caught.exception.field, "hotword_rows")
                self.assertEqual(str(caught.exception), message)
        path.unlink()  # Windows上仍被解析器占用的文件不能删除。
